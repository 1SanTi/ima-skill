# -*- coding: utf-8 -*-
"""
将 IMA 笔记 Markdown（政策汇编类）按文件分节整理，生成一个大 Word 文档。
用法: python build_docx.py <input.md> <output.docx>
"""
import sys, re, html
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

# ---------- 基础工具 ----------

def clean(s):
    s = s.replace('\u200b', '')
    s = re.sub(r'<br\s*/?>', '\n', s, flags=re.I)
    s = re.sub(r'<[^>]+>', '', s)          # 去掉 span/u/file 等所有标签
    s = html.unescape(s)
    for a, b in [('\\_', '_'), ('\\~', '~'), ('\\&', '&'), ('\\[', '['),
                 ('\\]', ']'), ('\\*', '*'), ('\\#', '#')]:
        s = s.replace(a, b)
    s = re.sub(r'(?<=[\u4e00-\u9fff])[ \t]+(?=[\u4e00-\u9fff])', '', s)  # 去掉汉字间空格
    return s

_SPLIT_RE = re.compile(
    r'(?=[一二三四五六七八九十]+、)|(?=（[一二三四五六七八九十]+）)|(?<=[\u4e00-\u9fff])(?=\d{1,2}\.)')

def expand_long(text, limit=500):
    """把被压成一整段的正文按 (一)/一、/1. 等标记切分为多行"""
    if len(text) <= limit:
        return [text]
    return [p for p in _SPLIT_RE.split(text) if p.strip()]

def set_run_font(run, cn='宋体', en=None, size=None, bold=None, color=None):
    run.font.name = en or cn
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), en or cn)
    rFonts.set(qn('w:hAnsi'), en or cn)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = RGBColor(*color)

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), r_id)
    run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '宋体'); rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); run.append(t)
    hl.append(run); paragraph._p.append(hl)

# ---------- 样式 ----------

def build_styles(doc):
    st = doc.styles['Normal']
    st.font.name = '宋体'; st.font.size = Pt(12)
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体'); rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5

    cfg = {'Heading 1': (16, (0, 0, 0)), 'Heading 2': (14, (0, 0, 0)),
           'Heading 3': (12.5, (0, 0, 0)), 'Heading 4': (12, (0, 0, 0))}
    for name, (sz, col) in cfg.items():
        s = doc.styles[name]
        s.font.name = '黑体'; s.font.size = Pt(sz); s.font.bold = True
        s.font.color.rgb = RGBColor(*col)
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
        rFonts.set(qn('w:eastAsia'), '黑体'); rFonts.set(qn('w:ascii'), 'Arial')
        rFonts.set(qn('w:hAnsi'), 'Arial')

def setup_page(doc):
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)

# ---------- 正文段落 ----------

BOLD_RE = re.compile(r'(\*\*[^*]+\*\*)')
LINK_RE = re.compile(r'\[([^\]]+)\]\((https?://[^)\s]+)\)')

def _is_h2(t): return bool(re.match(r'^[一二三四五六七八九十百]+、.{0,30}$', t))
def _is_h3(t): return bool(re.match(r'^（[一二三四五六七八九十]+）.{0,30}$', t))
def _is_article(t): return bool(re.match(r'^第[一二三四五六七八九十百]+[条章节]', t)) and len(t) < 40

def add_rich(p, text, base_bold=False):
    """把含 **加粗** 与 [文本](链接) 的文本写入段落 p"""
    pos = 0
    pattern = re.compile(r'\*\*([^*]+)\*\*|\[([^\]]+)\]\((https?://[^)\s]+)\)')
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            set_run_font(r, bold=base_bold or None)
        if m.group(1) is not None:
            r = p.add_run(m.group(1)); set_run_font(r, bold=True)
        else:
            add_hyperlink(p, m.group(3), m.group(2))
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:]); set_run_font(r, bold=base_bold or None)
    return p

def add_body_line(doc, line):
    line = line.strip()
    if not line or line == '***' or line == '---':
        return
    # markdown 标题
    m = re.match(r'^(#{1,6})\s+(.*)$', line)
    if m:
        lvl = min(len(m.group(1)) + 1, 4)  # 正文内 # -> H2
        doc.add_heading(clean(m.group(2)).strip(), level=lvl)
        return
    # 拆 bold 片段
    for tok in BOLD_RE.split(line):
        if not tok.strip():
            continue
        if tok.startswith('**') and tok.endswith('**'):
            inner = clean(tok[2:-2]).strip()
            if _is_h2(inner):
                doc.add_heading(inner, level=2)
            elif _is_h3(inner):
                doc.add_heading(inner, level=3)
            elif _is_article(inner):
                doc.add_heading(inner, level=4)
            else:
                p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(24)
                add_rich(p, inner, base_bold=True)
        else:
            inner = clean(tok).strip()
            if not inner:
                continue
            for piece in expand_long(inner):
                piece = piece.strip()
                if not piece:
                    continue
                if _is_h2(piece):
                    doc.add_heading(piece, level=2)
                elif _is_h3(piece):
                    doc.add_heading(piece, level=3)
                elif _is_article(piece):
                    doc.add_heading(piece, level=4)
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.first_line_indent = Pt(24)
                    add_rich(p, piece)

# ---------- 解析 ----------

DOC_TYPE_KW = ('通知', '意见', '方案', '规划', '纲要', '办法', '条例', '指南',
               '标准', '决定', '计划', '报告', '框架', '措施', '行动', '规则', '法')

HEAD_KW = ('通知', '意见', '方案', '规划', '纲要', '办法', '条例',
           '指南', '决定', '计划', '报告', '措施')

def is_file_heading(line):
    m = re.match(r'^(\d{1,2})\s*\.\s*(\S.*)$', line)
    if not m:
        return False
    rest = m.group(2)
    # 含 markdown 链接的条目行
    if re.search(r'\[[^\]]+\]\(https?://', rest):
        return True
    head = re.split(r'[。；：]', rest)[0].strip()
    if len(head) > 60:
        return False
    if '《' in head[:40]:
        return True
    if any(head.endswith(k) for k in HEAD_KW):
        return True
    return False

def parse_note(md):
    lines = md.split('\n')
    title = '未命名文档'
    for ln in lines:
        if ln.startswith('# '):
            title = clean(ln[2:]).strip(); break
    # 收集节起点
    idx = [i for i, ln in enumerate(lines)
           if is_file_heading(clean(ln).strip())]
    # 尾部区（参考文献 / 链接验证）
    tail = [i for i, ln in enumerate(lines) if re.match(r'^#{1,3}\s*(参考文献|链接验证)', clean(ln).strip())]
    sections = []
    for k, i in enumerate(idx):
        end = idx[k + 1] if k + 1 < len(idx) else len(lines)
        # 若该节跨越尾部标题，则截断
        for t in tail:
            if i < t < end:
                end = t
        raw_head = clean(lines[i]).strip()
        m = re.match(r'^(\d{1,2})\s*\.\s*(\S.*)$', raw_head)
        num, rest = m.group(1), m.group(2)
        lm = re.search(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', rest)
        url = lm.group(2) if lm else None
        t = re.sub(r'\[([^\]]+)\]\([^)]+\)', '', rest).rstrip('：: ').strip()
        t = t.strip('《》') if t.startswith('《') and t.endswith('》') else t
        body = lines[i + 1:end]
        sections.append({'num': num, 'title': t, 'url': url, 'body': body})
    tails = []
    for k, i in enumerate(tail):
        end = tail[k + 1] if k + 1 < len(tail) else len(lines)
        heads = [j for j in idx if j > i]
        if heads: end = min(end, heads[0])
        h = clean(lines[i]).strip().lstrip('#').strip()
        tails.append({'title': h, 'body': lines[i + 1:end]})
    return title, sections, tails

# ---------- 主流程 ----------

def main(inp, outp):
    md = open(inp, encoding='utf-8').read()
    title, sections, tails = parse_note(md)

    doc = Document()
    setup_page(doc); build_styles(doc)

    # 封面
    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_before = Pt(160)
    r = tp.add_run(title); set_run_font(r, cn='黑体', size=28, bold=True)
    sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sp.add_run('（政策文件汇编）'); set_run_font(r, cn='黑体', size=16)
    sp2 = doc.add_paragraph(); sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sp2.add_run(f'共收录 {len(sections)} 份文件'); set_run_font(r, cn='宋体', size=12)
    sp3 = doc.add_paragraph(); sp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sp3.add_run('整理自 IMA 笔记'); set_run_font(r, cn='宋体', size=11, color=(120, 120, 120))
    doc.add_page_break()

    # 文件清单
    doc.add_heading('收录文件清单', level=1)
    for s in sections:
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.4
        r = p.add_run(f"{s['num']}. {s['title']}"); set_run_font(r, cn='宋体', size=11)
    doc.add_page_break()

    # 各文件
    for si, s in enumerate(sections):
        if si > 0:
            doc.add_page_break()
        h = doc.add_heading(f"{s['num']}. {s['title']}", level=1)
        if s['url']:
            p = doc.add_paragraph()
            r = p.add_run('来源：'); set_run_font(r, cn='宋体', size=9, color=(120, 120, 120))
            add_hyperlink(p, s['url'], s['url'])
        for ln in s['body']:
            raw = clean(ln)
            if 'file cosKey' in ln or 'isAttachment' in ln:
                p = doc.add_paragraph()
                r = p.add_run('【本文件为 PDF 附件，详见笔记原附件】')
                set_run_font(r, cn='楷体', size=11, color=(150, 60, 60))
                continue
            for sub in raw.split('\n'):
                add_body_line(doc, sub)

    # 尾部
    for t in tails:
        doc.add_page_break()
        doc.add_heading(t['title'], level=1)
        for ln in t['body']:
            for sub in clean(ln).split('\n'):
                add_body_line(doc, sub)

    doc.save(outp)
    print(f'OK sections={len(sections)} tails={len(tails)} -> {outp}')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
