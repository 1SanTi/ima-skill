# -*- coding: utf-8 -*-
"""
docx 表格就地修改工具库（自包含）。

原则：只在原文件/原表格上改文本与对齐，不重建表格、不重排版，最大限度保留原格式。

库用法：
    from docx_table_reviser import open_doc, save_doc, rep, rep_cell, lc, rewrite_cell, set_valign

命令行自检：
    python3 docx_table_reviser.py <file.docx>
"""
import os
import sys
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH


def open_doc(path):
    return Document(path)


def save_doc(doc, out):
    """保存文档。中文路径先写英文临时文件再重命名，规避 zip 写入异常。"""
    if any(ord(c) > 127 for c in out):
        tmp = os.path.join(os.path.dirname(out) or ".", "_docx_tmp.docx")
        doc.save(tmp)
        if os.path.exists(out):
            os.remove(out)
        os.rename(tmp, out)
    else:
        doc.save(out)


def _set_para_text(p, text):
    """把整段文本压到第一个 run，保留该 run 的格式。"""
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def rep(p, old, new):
    """段内替换：先逐 run（保各 run 格式），跨 run 时整段兜底。返回是否替换。"""
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    full = p.text
    if old in full:
        _set_para_text(p, full.replace(old, new))
        return True
    return False


def rep_cell(cell, old, new):
    n = 0
    for p in cell.paragraphs:
        if rep(p, old, new):
            n += 1
    return n


def lc(row):
    """去重后的逻辑单元格（合并格只取一次）。"""
    seen = []
    for c in row.cells:
        if c._tc not in [s._tc for s in seen]:
            seen.append(c)
    return seen


def set_valign(cell, val="center"):
    """设置单元格垂直对齐：top / center / bottom。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn('w:vAlign')):
        tcPr.remove(e)
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), val)
    tcPr.append(va)


def rewrite_cell(cell, lines, align=None, valign=None):
    """
    保留首段/首 run 字体格式，重写整格文本。
    lines: 文本列表，每个元素生成一个段落（用于格内换行）。
    align: WD_ALIGN_PARAGRAPH.CENTER / LEFT / RIGHT（None=保持）。
    valign: 'center' 等，设置垂直对齐（None=不改）。
    """
    ps = cell.paragraphs
    base = ps[0]
    rpr = None
    for p in ps:
        for r in p.runs:
            if r._element.rPr is not None:
                rpr = deepcopy(r._element.rPr)
                break
        if rpr is not None:
            break
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    for r in list(base.runs):
        r._element.getparent().remove(r._element)
    if align is not None:
        base.alignment = align

    def fill(para, text):
        r = para.add_run(text)
        if rpr is not None:
            ex = r._element.find(qn('w:rPr'))
            if ex is not None:
                r._element.remove(ex)
            r._element.insert(0, deepcopy(rpr))

    fill(base, lines[0])
    last = base._p
    for tx in lines[1:]:
        np = deepcopy(base._p)
        for ch in list(np):
            if ch.tag != qn('w:pPr'):
                np.remove(ch)
        last.addnext(np)
        last = np
        fill(Paragraph(np, base._parent), tx)
    if valign:
        set_valign(cell, valign)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python3 docx_table_reviser.py <file.docx>")
        sys.exit(0)
    d = open_doc(sys.argv[1])
    print("段落数:", len(d.paragraphs), "表格数:", len(d.tables))
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            print("  P%d: %r" % (i, p.text))
    for ti, t in enumerate(d.tables):
        print("  表格%d: %d行 x %d列" % (ti, len(t.rows), len(t.columns)))
        for ri, row in enumerate(t.rows):
            cells = []
            for c in lc(row):
                cells.append(c.text.strip().replace("\n", "/"))
            print("    R%02d: %s" % (ri, " | ".join(cells)))
