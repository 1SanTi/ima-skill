# -*- coding: utf-8 -*-
"""
在原《通知》docx 上直接填写申报表，保持原表格式完全不变。
用法：
    python3 fill_form.py --template 通知.docx --data 候选人.json --out 申报表.docx
"""
import argparse
import json
import os
import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 候选人数据（示例：刘湘仔）
DEFAULT_DATA = {
    "姓名": "刘湘仔",
    "申报奖励名称": "衡东县优秀班主任",
    "工作单位全称": "衡东县蓬源镇中学",
    "填报时间": "2026年9月",
    "性别": "男",
    "民族": "汉",
    "出生日期": "1972年8月",
    "籍贯": "湖南省衡东县",
    "政治面貌": "群众",
    "身份证号码": "430424197208046214",
    "参加工作时间": "1995年9月",
    "学历学位": "本科",
    "教龄": "31年",
    "现任行政职务及级别": "",
    "任现行政职务时间": "",
    "专业技术职称": "中小学一级教师",
    "是否乡村教师_农村工作时长": "是，农村学校工作31年",
    "是否班主任_工作年限": "是，14年",
    "是否学前特教民办职教": "否",
    "是否学校班子成员_工作年限": "否",
    "教师资格证号": "20024304920001195",
    "工作单位": "衡东县蓬源镇中学",
    "联系电话": "13272306627",
    "工作简历": [["1995年9月—至今", "衡东县蓬源镇中学", "教师"]],
    "荣誉": [
        ["师德先进个人", "2020年12月", "蓬源教管中心"],
        ["年度优秀", "2020年12月", "县教育局"],
        ["县优秀班主任", "2021年12月", "县教育局"],
        ["年度优秀", "2021年11月", "县政府"],
        ["县优秀班主任", "2023年5月", "县教育局"],
        ["县优秀乡村教师", "2024年9月", "县教育局"],
    ],
    "教学工作量": [
        "蓬源中学从事九年级数学、八年级生物、九年级体育（每周12节课）",
        "蓬源中学教九年级数学、八年级生物（每周14节课）",
        "蓬源中学教九年级数学、八年级生物（每周13节课）",
        "蓬源中学教九年级数学、八年级生物（每周11节课）",
        "蓬源中学教九年级数学、八年级生物（每周12节课）",
    ],
    "先进事迹": "刘湘仔同志扎根乡村教育三十一年，长期担任班主任并坚守教学一线。他恪尽职守、爱生如子，每天清晨带领学生晨跑，以高度的责任心和无私奉献精神默默耕耘。近年因青年骨干教师外流，他主动承担初三数学、初三体育和初二生物三门中考科目的教学任务，顾全大局、勇挑重担，精心备课、认真上课、个别辅导、科学训练，所带班级中考成绩逐年攀升，2024年、2025年分别在全县综合排名中位列第三、第四名。他连续担任七届班主任，班级管理规范有序、班风学风良好。2018年中考前夕，他因胃出血住院，病情稍有好转便提前出院，带病坚持工作，主动提出双休日为学生补课、不计报酬，所带班级当年中考综合排名较前一年前进十一名，一举改变学校落后局面。他发挥体育特长，主动承担学校体育训练任务，多次带领学生在全县中小学生运动会中取得团体总分第一名，个人多次获“县优秀教练员”称号。他以“吃得苦、霸得蛮、耐得烦”的精神忘我工作，深受学生爱戴和家长信赖，是乡村教育战线上的优秀典型。",
}


def logic_cells(row):
    """返回去重后的逻辑单元格列表（合并单元格只取一次）"""
    seen = []
    for c in row.cells:
        if c._tc not in [s._tc for s in seen]:
            seen.append(c)
    return seen


def fill_cell(cell, text, align="center", font="FangSong", size=12):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = None
    run = p.add_run(text)
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    return cell


def fill_cover(p, text):
    """把封面下划线空位替换为填写值"""
    for r in p.runs:
        rpr = r._element.rPr
        if rpr is not None and rpr.find(qn("w:u")) is not None:
            r.text = text
            return True
    return False


def strip_space(s):
    return s.replace(" ", "").replace("\u3000", "").replace("\t", "").replace("\n", "")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, help="原始通知 docx 路径")
    ap.add_argument("--data", default=None, help="候选人 JSON 路径")
    ap.add_argument("--out", required=True, help="输出申报表 docx 路径")
    args = ap.parse_args()

    data = DEFAULT_DATA
    if args.data:
        with open(args.data, "r", encoding="utf-8") as f:
            data = json.load(f)

    # 复制模板，避免改动原始文件
    shutil.copy(args.template, args.out)
    doc = Document(args.out)

    # ---- 先保存申报表三个表格的引用（删除其他附件后 doc.tables 索引会变）----
    info_tbl = doc.tables[2]    # 基本信息 + 工作简历 + 荣誉
    work_tbl = doc.tables[3]    # 教学工作量 + 先进事迹
    opinion_tbl = doc.tables[4]  # 本人意见/学校意见/审批意见

    # ---- 1. 删除非申报表（附件3）内容 ----
    body = doc.element.body
    title_el = None
    for p in doc.paragraphs:
        if strip_space(p.text) == "衡东县教育系统先进人物申报表":
            title_el = p._p
            break
    if title_el is None:
        raise SystemExit("未找到申报表标题段落")
    tbl4_el = opinion_tbl._tbl  # 申报表最后一个表（意见/审批表）

    children = list(body.iterchildren())
    title_idx = children.index(title_el)
    tbl4_idx = children.index(tbl4_el)
    for i, child in enumerate(children):
        if child.tag in (qn("w:p"), qn("w:tbl")):
            if i < title_idx or i > tbl4_idx:
                body.remove(child)

    # ---- 2. 填写封面 ----
    for p in doc.paragraphs:
        key = strip_space(p.text)
        if key == "姓名":
            fill_cover(p, data["姓名"])
        elif key == "申报奖励名称":
            fill_cover(p, data["申报奖励名称"])
        elif key == "工作单位（全称）":
            fill_cover(p, data["工作单位全称"])
        elif key == "填报时间":
            fill_cover(p, data["填报时间"])

    # ---- 3. 填写表格2：基本信息 + 工作简历 + 荣誉 ----
    t2 = info_tbl
    base_info = [
        (data["姓名"], data["性别"]),
        (data["民族"], data["出生日期"]),
        (data["籍贯"], data["政治面貌"]),
        (data["身份证号码"], data["参加工作时间"]),
        (data["学历学位"], data["教龄"]),
        (data["现任行政职务及级别"], data["任现行政职务时间"]),
        (data["专业技术职称"], data["是否乡村教师_农村工作时长"]),
        (data["是否班主任_工作年限"], data["是否学前特教民办职教"]),
        (data["是否学校班子成员_工作年限"], data["教师资格证号"]),
        (data["工作单位"], data["联系电话"]),
    ]
    for r in range(10):
        lc = logic_cells(t2.rows[r])
        fill_cell(lc[1], base_info[r][0], align="center")
        fill_cell(lc[3], base_info[r][1], align="center")

    # 工作简历（行11-15，逻辑格1=时间,2=所在单位,3=从事工作）
    for i, item in enumerate(data["工作简历"][:5]):
        lc = logic_cells(t2.rows[11 + i])
        fill_cell(lc[1], item[0], align="center")
        fill_cell(lc[2], item[1], align="center")
        fill_cell(lc[3], item[2], align="center")

    # 荣誉（行17-22，逻辑格1=获奖名称,2=获奖时间,3=授予单位）
    for i, item in enumerate(data["荣誉"][:6]):
        lc = logic_cells(t2.rows[17 + i])
        fill_cell(lc[1], item[0], align="center")
        fill_cell(lc[2], item[1], align="center")
        fill_cell(lc[3], item[2], align="center")

    # ---- 4. 填写表格3：教学工作量 + 先进事迹 ----
    t3 = work_tbl
    for i, content in enumerate(data["教学工作量"][:5]):
        lc = logic_cells(t3.rows[1 + i])
        fill_cell(lc[2], content, align="left")
    # 先进事迹（行6，逻辑格1）
    lc6 = logic_cells(t3.rows[6])
    fill_cell(lc6[1], data["先进事迹"], align="left")

    # 中文路径写入的兼容处理：先存英文临时文件再重命名，避免 zip 写入异常
    out = args.out
    if any(ord(c) > 127 for c in out):
        tmp = os.path.join(os.path.dirname(out) or ".", "_fill_form_tmp.docx")
        doc.save(tmp)
        if os.path.exists(out):
            os.remove(out)
        os.rename(tmp, out)
    else:
        doc.save(out)
    print(f"已生成：{out}")


if __name__ == "__main__":
    main()
