# -*- coding: utf-8 -*-
"""研修总结 Word 生成器（JSON 驱动）。

用法：
    python3 build_summary_docx.py input.json -o out.docx

输入 JSON：
{
  "title": "文档大标题（可含 \\n 换行）",
  "blocks": [
    {"style": "h1",   "text": "一、研修概况"},
    {"style": "h2",   "text": "（一）……"},
    {"style": "para", "text": "正文段落……"},
    {"style": "sign", "text": "总结人：某某"}
  ]
}
style 取值：title / h1 / h2 / para / sign（sign = 右对齐、无首行缩进的署名行）
排版规范与 ima-doc 一致：A4；正文宋体+Times New Roman 12pt、1.5 倍行距、首行缩进 2 字符；标题黑体。
"""
import argparse
import json
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_style_cjk_font(style, cn_font, en_font):
    style.font.name = en_font
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn_font)
    rFonts.set(qn('w:ascii'), en_font)
    rFonts.set(qn('w:hAnsi'), en_font)


def clean_heading_style(style):
    style.font.color.rgb = RGBColor(0, 0, 0)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        for attr in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
            if rFonts.get(qn(attr)):
                del rFonts.attrib[qn(attr)]
    pPr = style.element.find(qn('w:pPr'))
    if pPr is not None:
        pBdr = pPr.find(qn('w:pBdr'))
        if pBdr is not None:
            pPr.remove(pBdr)


def build(data, out_path):
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    set_style_cjk_font(normal, '宋体', 'Times New Roman')
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)

    title_style = doc.styles['Title']
    set_style_cjk_font(title_style, '黑体', 'Arial')
    clean_heading_style(title_style)
    title_style.font.size = Pt(18)

    for i in range(1, 4):
        clean_heading_style(doc.styles[f'Heading {i}'])
        set_style_cjk_font(doc.styles[f'Heading {i}'], '黑体', 'Arial')
    h1 = doc.styles['Heading 1']
    h1.font.size = Pt(15)
    h1.paragraph_format.first_line_indent = Pt(0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h2 = doc.styles['Heading 2']
    h2.font.size = Pt(13)
    h2.paragraph_format.first_line_indent = Pt(0)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)

    title = data.get('title', '')
    if title:
        t = doc.add_heading(title, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in t.runs:
            r.font.size = Pt(18)

    for b in data.get('blocks', []):
        style = b.get('style', 'para')
        text = b.get('text', '')
        if style == 'title':
            h = doc.add_heading(text, level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in h.runs:
                r.font.size = Pt(18)
        elif style == 'h1':
            doc.add_heading(text, level=1)
        elif style == 'h2':
            doc.add_heading(text, level=2)
        elif style == 'sign':
            p = doc.add_paragraph(text)
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:  # para
            doc.add_paragraph(text)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description='研修总结 Word 生成器')
    ap.add_argument('input', help='输入 JSON 文件路径')
    ap.add_argument('-o', '--out', required=True, help='输出 .docx 路径')
    args = ap.parse_args()

    with open(args.input, 'r', encoding='utf-8') as f:
        data = json.load(f)
    out = build(data, args.out)
    print('saved:', out)


if __name__ == '__main__':
    main()
