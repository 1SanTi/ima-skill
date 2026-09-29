# -*- coding: utf-8 -*-
"""数据驱动：按 data.json 填写“公费师范生入编三表”并可选打包。

用法：
  python3 fill_entry_forms.py --templates ./templates --data ./data.json --out ./out [--photo photo.jpg] [--zip]

templates 目录下（前缀匹配即可）：
  3.*.xlsx  -> 《机关事业单位人员编制审批表》（用“空白表”sheet）
  4.*.docx  -> 《新进编人员登记表》
  5.*.docx  -> 《聘用人员花名册》

data.json 字段（缺项用空字符串 "" 或省略，程序会留空）：
{
  "seq": 43, "name": "李泽宇", "sex": "男", "nation": "汉族",
  "birth": "2004.01", "idno": "430000200401010000",
  "jiguan": "湖南云溪", "hukou": "云溪县",
  "hukou_type": "农业户口",        // 或 "非农户口"，空则不勾
  "marital": "未婚",               // 或 "已婚"，空则不勾
  "zzmm": "共青团员",
  "edu": "本科", "degree": "学士",
  "school": "衡阳师范学院", "major": "应用心理学", "gradtime": "2026年6月",
  "unit": "云溪县第六中学", "unit_nature": "全额事业",
  "title": "",                     // 职称，空则留空
  "qualification": "初中心理教师资格证", "qual_time": "",
  "teaching_subject": "初中心理教育",
  "work_start": "",                // 参加工作时间
  "entry_method": "⑩其他",          // 使用编制进入方式，勾选项文字
  "resume": "2019.09—2022.06  ... 学生\\n2022.09—2026.06  ... 学生",
  "family": [ {"relation":"父亲","name":"罗仁成","age":46,"political":"群众","unit":"经商"},
              {"relation":"母亲","name":"周玉良","age":44,"political":"群众","unit":"务农"} ],
  "contact_person": "郭志强", "contact_phone": "13800000002",
  "date": "2026年9月1日",
  "exam": {"written":"","interview":"","total":""}, "physical":"", "assess":"",
  "hiring_opinion_unit": "同意聘用", "hiring_opinion_dept": "同意聘用",
  "hiring_opinion_hr": "同意从2026年9月1日聘起",
  "reason": "2026年公费师范生毕业分配"
}
"""
import os, glob, json, argparse, zipfile
import openpyxl
from docx import Document
from docx.shared import Cm
from form_fill_utils import set_docx_cell, grid_cell, insert_photo_fit, xlsx_insert_photo


def find(prefix_dir, pat):
    hits = glob.glob(os.path.join(prefix_dir, pat))
    if not hits:
        raise FileNotFoundError(f'{pat} not found in {prefix_dir}')
    return sorted(hits)[0]


def mark(txt, chosen):
    """把选项串里的 □ 勾成 √：chosen 为要勾的选项关键字。"""
    out = txt
    for key in ['招录', '招聘', '调任', '选调', '选拔', '调入', '政策性安置', '人才引进', '内部调整', '其他']:
        tag = '√' if key in (chosen or '') else '□'
        # 替换“关键词”后紧跟的方框
        i = out.find(key)
        while i != -1:
            j = i + len(key)
            while j < len(out) and out[j] in ' 　':
                j += 1
            if j < len(out) and out[j] in '□√':
                out = out[:j] + tag + out[j + 1:]
            i = out.find(key, j)
    return out


def fill_xlsx(path, d, photo, out):
    wb = openpyxl.load_workbook(path)
    ws = wb['空白表'] if '空白表' in wb.sheetnames else wb.active
    ws['B3'] = d.get('unit', '')
    ws['E3'] = d.get('unit_nature', '')
    if d.get('entry_method'):
        ws['B4'] = mark(ws['B4'].value, d['entry_method'])
    ws['C5'] = d.get('name', ''); ws['E5'] = d.get('sex', ''); ws['G5'] = d.get('birth', '')
    ws['C6'] = d.get('idno', ''); ws['E6'] = d.get('nation', ''); ws['G6'] = d.get('jiguan', '')
    ws['C7'] = d.get('work_start', ''); ws['E7'] = d.get('zzmm', ''); ws['G7'] = d.get('edu', '')
    ws['C8'] = d.get('major', ''); ws['E8'] = d.get('title', ''); ws['G8'] = d.get('org_prev', '无')
    ws['B9'] = d.get('resume', '')
    ws['B9'].alignment = openpyxl.styles.Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws['C6'].alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
    ws['B15'] = d.get('reason', '')
    ws['A21'] = '同意'; ws['A23'] = d.get('date', '')
    ws['D21'] = '同意'; ws['D23'] = d.get('date', '')
    if photo and os.path.exists(photo) and ws['H5'].value is not None:
        ws['H5'] = ''          # 清“近期一寸彩色照片”占位
        xlsx_insert_photo(ws, 'H5', photo, max_w_cm=2.7, max_h_cm=4.0)
    if '模板' in wb.sheetnames:
        del wb['模板']
    o = os.path.join(out, '1.机关事业单位人员编制审批表.docx'.replace('.docx', '.xlsx'))
    os.makedirs(out, exist_ok=True)
    wb.save(o)
    return o


def fill_dengji(path, d, photo, out):
    doc = Document(path)
    t = doc.tables[0]
    vals = {(0, 1): d.get('name', ''), (0, 5): d.get('sex', ''), (0, 8): d.get('nation', ''),
            (1, 1): d.get('birth', ''), (1, 5): d.get('zzmm', ''), (1, 8): d.get('jiguan', ''),
            (2, 1): d.get('hukou', ''),
            (2, 5): '非农户口□  农业户口√' if d.get('hukou_type') == '农业户口' else
                    ('非农户口√  农业户口□' if d.get('hukou_type') == '非农户口' else '非农户口□  农业户口□'),
            (2, 8): '已婚□  未婚√' if d.get('marital') == '未婚' else
                    ('已婚√  未婚□' if d.get('marital') == '已婚' else '已婚□  未婚□'),
            (3, 1): d.get('idno', ''),
            (4, 1): d.get('edu', ''), (4, 5): d.get('degree', ''),
            (4, 9): f"{d.get('school','')}  {d.get('major','')}  {d.get('gradtime','')}".strip(),
            (5, 4): d.get('title', ''), (6, 4): d.get('qualification', ''), (6, 11): d.get('qual_time', ''),
            (7, 2): '无', (8, 2): d.get('unit', ''), (8, 9): d.get('unit_nature', ''),
            (9, 2): '管理岗位□\t\t专业技术岗位√ \t工勤技能岗位□', (10, 2): d.get('resume', '')}
    for (ri, gc), txt in vals.items():
        if txt:
            set_docx_cell(grid_cell(t.rows[ri], gc), txt, 'left' if (ri, gc) in [(4, 9), (7, 2), (10, 2)] else 'center')
    if photo and os.path.exists(photo):
        set_docx_cell(grid_cell(t.rows[0], 10), ' ' * 1)  # 清空
        c = grid_cell(t.rows[0], 10)
        for p in c.paragraphs[1:]:
            p._element.getparent().remove(p._element)
        p = c.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.add_run().add_picture(photo, width=Cm(2.3), height=Cm(3.45))
    t1 = doc.tables[1]
    for i, m in enumerate(d.get('family', [])[:4]):
        ri = 1 + i
        for gc, key in [(1, 'relation'), (2, 'name'), (3, 'age'), (4, 'political'), (5, 'unit')]:
            set_docx_cell(grid_cell(t1.rows[ri], gc), str(m.get(key, '')), 'center')
    if d.get('teaching_subject'):
        set_docx_cell(grid_cell(t1.rows[15], 2), f"任教学科：{d['teaching_subject']}", 'left')
    if d.get('hiring_opinion_unit'):
        set_docx_cell(grid_cell(t1.rows[13], 1), d['hiring_opinion_unit'], 'center')
    if d.get('hiring_opinion_dept'):
        set_docx_cell(grid_cell(t1.rows[13], 5), d['hiring_opinion_dept'], 'center')
    if d.get('hiring_opinion_hr'):
        set_docx_cell(grid_cell(t1.rows[14], 1), d['hiring_opinion_hr'], 'center')
    o = os.path.join(out, '2.新进编人员登记表.docx')
    doc.save(o)
    return o


def fill_huaming(path, d, out):
    doc = Document(path)
    for p in doc.paragraphs:
        if '招聘单位' in p.text:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            r = p.add_run(f"招聘单位：{d.get('unit','')}        联系人：{d.get('contact_person','')}        "
                          f"联系电话：{d.get('contact_phone','')}        {d.get('date','')}")
            from form_fill_utils import set_font
            set_font(r)
            break
    t = doc.tables[0]
    row = t.rows[2]
    data = [str(d.get('seq', '')), d.get('name', ''), d.get('sex', ''), d.get('birth', ''), d.get('zzmm', ''),
            d.get('edu', ''), d.get('degree', ''), d.get('school', ''), d.get('major', ''),
            d.get('qualification', ''), '专业技术岗位',
            f"任教学科：{d['teaching_subject']}" if d.get('teaching_subject') else '']
    for ci, v in enumerate(data):
        set_docx_cell(row.cells[ci], v, 'center')
    # 单人：删多余空数据行，保持 1 页
    for idx in range(len(t.rows) - 3, 2, -1):
        tr = t.rows[idx]._tr
        if all(not c.text.strip() for c in t.rows[idx].cells):
            tr.getparent().remove(tr)
    o = os.path.join(out, '3.聘用人员花名册.docx')
    doc.save(o)
    return o


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--templates', required=True)
    ap.add_argument('--data', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--photo', default='')
    ap.add_argument('--zip', action='store_true')
    a = ap.parse_args()
    d = json.load(open(a.data, encoding='utf-8'))
    os.makedirs(a.out, exist_ok=True)
    outs = [fill_xlsx(find(a.templates, '3.*.xlsx'), d, a.photo, a.out),
            fill_dengji(find(a.templates, '4.*.docx'), d, a.photo, a.out),
            fill_huaming(find(a.templates, '5.*.docx'), d, a.out)]
    print('generated:', *outs, sep='\n  ')
    if a.zip:
        z = os.path.join(a.out, f"{d.get('name','入编')}入编材料.zip")
        with zipfile.ZipFile(z, 'w', zipfile.ZIP_DEFLATED) as zf:
            for f in outs:
                zf.write(f, os.path.basename(f))
        print('zip:', z)


if __name__ == '__main__':
    main()
