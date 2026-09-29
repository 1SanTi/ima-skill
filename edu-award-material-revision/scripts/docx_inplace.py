# -*- coding: utf-8 -*-
"""
docx 原表就地修改工具库。

原则：只在原文件/原表格上修改文本、删除行，不重建表格，最大限度保留原格式。

可作为库导入：
    from docx_inplace import open_doc, save_doc, rep, rep_all, rep_cell, del_row, lc

也可命令行自检：
    python3 docx_inplace.py <file.docx>
"""
import os
import sys
from docx import Document


def open_doc(path):
    return Document(path)


def save_doc(doc, out):
    """保存文档。中文路径先写英文临时文件再重命名，规避某些环境 zip 写入异常。"""
    if any(ord(c) > 127 for c in out):
        tmp = os.path.join(os.path.dirname(out) or ".", "_docx_inplace_tmp.docx")
        doc.save(tmp)
        if os.path.exists(out):
            os.remove(out)
        os.rename(tmp, out)
    else:
        doc.save(out)


def _set_para_text(p, text):
    """把整段文本压缩到第一个 run，保留该 run 的格式。"""
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def rep(p, old, new):
    """
    在段落 p 内把 old 替换为 new。
    先尝试逐 run 替换（不破坏各 run 格式）；若 old 跨多个 run（常见于 Word 把一句拆成多 run），
    则退化为整段替换（保留首 run 格式）。返回是否发生替换。
    """
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    full = p.text
    if old in full:
        _set_para_text(p, full.replace(old, new))
        return True
    return False


def rep_cell(cell, old, new):
    n = 0
    for p in cell.paragraphs:
        if rep(p, old, new):
            n += 1
    return n


def rep_all(doc, old, new):
    """对文档所有段落与表格单元格做替换。返回替换次数。"""
    n = 0
    for p in doc.paragraphs:
        if rep(p, old, new):
            n += 1
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                n += rep_cell(c, old, new)
    return n


def lc(row):
    """返回去重后的逻辑单元格（合并单元格只取一次），用于按逻辑列定位。"""
    seen = []
    for c in row.cells:
        if c._tc not in [s._tc for s in seen]:
            seen.append(c)
    return seen


def del_row(row):
    """删除表格行（在 xml 层移除 w:tr）。"""
    row._tr.getparent().remove(row._tr)


def find_para(doc, contains):
    """返回包含指定文字的段落列表。"""
    return [p for p in doc.paragraphs if contains in p.text]


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python3 docx_inplace.py <file.docx>")
        sys.exit(0)
    d = open_doc(sys.argv[1])
    print("段落数:", len(d.paragraphs), "表格数:", len(d.tables))
    for i, t in enumerate(d.tables):
        print(f"  表格{i}: {len(t.rows)}行 x {len(t.columns)}列")
