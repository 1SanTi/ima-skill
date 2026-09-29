# -*- coding: utf-8 -*-
"""根据 questions.json 生成规范排版的 Word 问卷（.docx）。

用法: python3 spec_to_word.py questions.json 输出.docx
"""
import json, sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG = "宋体"
HEI = "黑体"
KAI = "楷体"


def set_font(run, name=SONG, size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), name)


def add_para(doc, text="", font=SONG, size=12, bold=False, align=None,
             space_before=0, space_after=4, line=1.5, color=None, indent=None,
             first_line_indent=None, keep_with_next=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if keep_with_next:
        pf.keep_with_next = True
        pf.keep_together = True
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    if text:
        run = p.add_run(text)
        set_font(run, font, size, bold, color)
    return p


def bottom_border(p, size=6):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '808080')
    pbdr.append(bottom)
    pPr.append(pbdr)


def build(spec, out_path):
    doc = Document()

    # 页面与默认字体
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)

    normal = doc.styles['Normal']
    normal.font.name = SONG
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

    # 标题
    add_para(doc, spec.get("title", "调查问卷"), font=HEI, size=20, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=4, line=1.4)
    if spec.get("subtitle"):
        add_para(doc, spec["subtitle"], font=KAI, size=13,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10, line=1.3)

    # 卷首语
    for i, line in enumerate(spec.get("intro", "").split("\n")):
        add_para(doc, line, font=KAI, size=11.5, line=1.5, space_after=2,
                 first_line_indent=0.9 if line else None)

    # 题号
    num = 0
    for section in spec["sections"]:
        add_para(doc, section["title"], font=HEI, size=14, bold=True,
                 space_before=10, space_after=6, line=1.4, keep_with_next=True)
        for q in section["questions"]:
            qtype = q["type"]
            if qtype == "text":
                num += 1
                mark = f"{num}."
                add_para(doc, f"{mark} {q['title']}", font=SONG, size=12, bold=True,
                         space_before=6, space_after=4, line=1.5, keep_with_next=True)
                # 答题横线
                for _ in range(3):
                    p = add_para(doc, "", size=12, space_after=10, line=1.6)
                    bottom_border(p)
                continue

            if qtype == "rating":
                num += 1
                add_para(doc, f"{num}. {q['title']}", font=SONG, size=12, bold=True,
                         space_before=6, space_after=4, line=1.5, keep_with_next=True)
                lo = q.get("min", 1)
                hi = q.get("max", 10)
                line_txt = "".join(f"○{i}　" for i in range(lo, hi + 1))
                add_para(doc, line_txt, font=SONG, size=12, space_after=2, line=1.5,
                         keep_with_next=True)
                if q.get("low_label") or q.get("high_label"):
                    add_para(doc, f"（{lo}={q.get('low_label', '')}，{hi}={q.get('high_label', '')}）",
                             font=KAI, size=10.5, space_after=2, line=1.4)
                continue

            num += 1
            add_para(doc, f"{num}. {q['title']}", font=SONG, size=12, bold=True,
                     space_before=6, space_after=4, line=1.5, keep_with_next=True)

            if qtype in ("single", "multi", "scale"):
                if qtype == "scale" or q.get("options_scale"):
                    opts = spec.get("scale_options", [])
                    box = "○"
                else:
                    opts = q.get("options", [])
                    box = "○" if qtype == "single" else "□"
                if box == "□" and len(opts) > 8:
                    # 多选项较多时每行 3 个，更整齐
                    line_txt = ""
                    for i, o in enumerate(opts):
                        line_txt += f"{box}{o}　　"
                        if (i + 1) % 3 == 0:
                            add_para(doc, line_txt, font=SONG, size=12,
                                     space_after=2, line=1.4)
                            line_txt = ""
                    if line_txt:
                        add_para(doc, line_txt, font=SONG, size=12,
                                 space_after=2, line=1.4)
                else:
                    line_txt = "".join(f"{box}{o}　　" for o in opts)
                    add_para(doc, line_txt, font=SONG, size=12,
                             space_after=2, line=1.5)

                if q.get("allow_other"):
                    add_para(doc, "□其他：＿＿＿＿＿＿＿＿", font=SONG, size=12,
                             space_after=2, line=1.5)

    add_para(doc, "", space_after=6)
    add_para(doc, spec.get("footer", "问卷到此结束，感谢您的参与！"),
             font=KAI, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    spec_path = sys.argv[1]
    out_path = sys.argv[2]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    build(spec, out_path)
    print("生成成功:", out_path)
